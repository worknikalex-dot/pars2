# -*- coding: utf-8 -*-
"""
Главный скрипт поиска цен на 24lek.ru.

Ищет цены:
1. В ближайших физических аптеках.
2. В фиксированном списке интернет-аптек.
3. Повторно ищет ненайденные цены физических аптек
   с фильтром по району.

Запуск:
    python main.py
"""

import sys
import time

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

import config
from normalizer import normalize_name
from geocoder import geocode_address, GeocodeError
from pharmacies import (
    load_pharmacies,
    pharmacies_in_radius,
    pharmacy_header,
)
from lek24_parser import search_drug_prices, _make_session


INTERNET_PHARMACY_IDS = [
    1105,  # Аптека.ру
    731,   # Аптека Легко
    897,   # Губернские аптеки
    1151,  # Фармэконом
    344,   # Нейрон
    813,   # ТаблеткиТут
]


# Соответствие названий районов параметру raon сайта 24lek.ru.
DISTRICT_RAON_IDS = {
    "железнодорожный": 1,
    "кировский": 2,
    "ленинский": 3,
    "октябрьский": 4,
    "свердловский": 5,
    "советский": 6,
    "центральный": 7,
}


def normalize_id(value):
    """Приводит ID аптеки к целому числу, если возможно."""
    try:
        return int(value)
    except (TypeError, ValueError):
        return value


def normalize_district(value):
    """Нормализует название района для сопоставления."""
    if value is None:
        return ""

    district = " ".join(str(value).strip().lower().split())
    district = district.removesuffix(" район").strip()

    return district


def get_raon_id(pharmacy):
    """Возвращает ID района аптеки либо None, если район неизвестен."""
    district = normalize_district(pharmacy.get("district"))
    return DISTRICT_RAON_IDS.get(district)


def read_products(path: str) -> list:
    """Читает названия препаратов из первого непустого столбца."""
    wb = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=True,
    )
    ws = wb.active
    names = []

    try:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is None:
                    continue

                text = str(cell).strip()
                if text:
                    names.append(text)
                break
    finally:
        wb.close()

    return names


def build_report(
    raw_products: list,
    nearby_pharmacies: list,
    internet_pharmacies: list,
    results: dict,
    out_path: str,
):
    """Создаёт Excel-отчёт с физическими и интернет-аптеками."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчёт"

    physical_headers = [
        pharmacy_header(ph)
        for ph in nearby_pharmacies
    ]

    internet_headers = [
        f"ИНТЕРНЕТ | {ph['name']} | {ph['address']}"
        for ph in internet_pharmacies
    ]

    headers = (
        ["Препарат (из списка)"]
        + physical_headers
        + internet_headers
    )

    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color="2F5496",
        end_color="2F5496",
        fill_type="solid",
    )
    internet_fill = PatternFill(
        start_color="548235",
        end_color="548235",
        fill_type="solid",
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.alignment = Alignment(
            wrap_text=True,
            vertical="center",
            horizontal="center",
        )

        if col_idx > 1 + len(nearby_pharmacies):
            cell.fill = internet_fill
        else:
            cell.fill = header_fill

    for raw_name in raw_products:
        price_map = results.get(raw_name, {})
        row = [raw_name]

        for ph in nearby_pharmacies:
            aid = normalize_id(ph["id"])
            price = price_map.get(aid)
            row.append(price if price is not None else "")

        for ph in internet_pharmacies:
            aid = normalize_id(ph["id"])
            price = price_map.get(aid)
            row.append(price if price is not None else "")

        ws.append(row)

    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 48

    for col_idx in range(2, len(headers) + 1):
        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = 28

    ws.row_dimensions[1].height = 75

    for row in ws.iter_rows(min_row=2, min_col=2):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00'

    wb.save(out_path)


def update_price_map(price_map, items, allowed_ids):
    """Добавляет найденные цены, сохраняя минимальную цену на аптеку."""
    for item in items:
        aid = normalize_id(item.get("apteka_id"))

        if aid not in allowed_ids:
            continue

        price = item.get("price")

        if price is None:
            continue

        if aid not in price_map or price < price_map[aid]:
            price_map[aid] = price


def main():
    print("=" * 70)
    print(" Поиск цен на лекарства на 24lek.ru")
    print(" Ближайшие аптеки + интернет-аптеки")
    print("=" * 70)

    address = input(
        "\nВведите адрес (например: Свободный 60): "
    ).strip()

    if not address:
        print("Адрес не введён. Завершение.")
        sys.exit(1)

    print("\n[1/5] Определяем координаты адреса через Dadata...")

    try:
        lat, lon, full_address = geocode_address(address)
    except GeocodeError as e:
        print(f"ОШИБКА геокодирования: {e}")
        sys.exit(1)

    print(f"   Адрес распознан как: {full_address}")
    print(f"   Координаты: {lat:.6f}, {lon:.6f}")

    print(
        f"\n[2/5] Загружаем {config.APTEKI_FILE} "
        f"и определяем аптеки..."
    )

    try:
        pharmacies = load_pharmacies()
    except FileNotFoundError:
        print(
            f"ОШИБКА: файл {config.APTEKI_FILE} "
            "не найден рядом со скриптом."
        )
        sys.exit(1)

    internet_ids = {
        normalize_id(aid)
        for aid in INTERNET_PHARMACY_IDS
    }

    pharmacies_by_id = {
        normalize_id(ph["id"]): ph
        for ph in pharmacies
    }

    internet_pharmacies = []

    for aid in INTERNET_PHARMACY_IDS:
        ph = pharmacies_by_id.get(normalize_id(aid))

        if ph is None:
            print(
                f"ПРЕДУПРЕЖДЕНИЕ: интернет-аптека с ID {aid} "
                "не найдена в apteki.json."
            )
            continue

        internet_pharmacies.append(ph)

    nearby_all = pharmacies_in_radius(
        pharmacies,
        lat,
        lon,
    )

    nearby = [
        ph for ph in nearby_all
        if normalize_id(ph["id"]) not in internet_ids
    ]

    print(f"   Ближайших физических аптек: {len(nearby)}")
    print(
        f"   Интернет-аптек из списка найдено: "
        f"{len(internet_pharmacies)}"
    )

    for ph in nearby:
        print(
            f"     - {ph['name']} — {ph['address']} "
            f"({ph['distance_m']} м) [id={ph['id']}]"
        )

    print("\n   Интернет-аптеки:")
    for ph in internet_pharmacies:
        print(
            f"     - {ph['name']} — {ph['address']} "
            f"[id={ph['id']}]"
        )

    needed_ids = {
        normalize_id(ph["id"])
        for ph in nearby + internet_pharmacies
    }

    if not needed_ids:
        print(
            "Не найдено аптек для поиска цен. "
            "Проверьте apteki.json."
        )
        sys.exit(1)

    print(
        f"\n[3/5] Загружаем список препаратов "
        f"из {config.PRODUCTS_FILE}..."
    )

    try:
        raw_products = read_products(config.PRODUCTS_FILE)
    except FileNotFoundError:
        print(
            f"ОШИБКА: файл {config.PRODUCTS_FILE} "
            "не найден рядом со скриптом."
        )
        sys.exit(1)

    print(f"   Загружено препаратов: {len(raw_products)}")

    print(
        "\n[4/5] Ищем цены для ближайших и интернет-аптек..."
    )

    physical_ids = {
        normalize_id(ph["id"])
        for ph in nearby
    }

    # Группируем физические аптеки по району.
    pharmacies_by_raon = {}

    for ph in nearby:
        aid = normalize_id(ph["id"])
        raon_id = get_raon_id(ph)

        if raon_id is None:
            print(
                f"[Районный поиск] Неизвестен район аптеки "
                f"{ph['name']} [id={aid}]: "
                f"{ph.get('district')!r}. "
                "Резервный поиск для неё пропускается."
            )
            continue

        pharmacies_by_raon.setdefault(
            raon_id, set()
        ).add(aid)

    session = _make_session()
    results = {}

    try:
        for idx, raw_name in enumerate(raw_products, start=1):
            norm = normalize_name(raw_name)

            if not norm:
                print(
                    f"   [{idx}/{len(raw_products)}] "
                    f"Пропуск: не удалось нормализовать "
                    f"{raw_name!r}"
                )
                results[raw_name] = {}
                continue

            print(
                f"\n   [{idx}/{len(raw_products)}] "
                f"{raw_name!r} -> запрос: '{norm}'"
            )

            # Этап 1. Основной поиск по всем районам.
            try:
                items = search_drug_prices(
                    norm,
                    needed_apteka_ids=needed_ids,
                    session=session,
                    raon_id=0,
                )
            except Exception as e:
                print(
                    f"      ОШИБКА основного поиска: {e}"
                )
                items = []

            price_map = {}
            update_price_map(price_map, items, needed_ids)

            # Этап 2. Резервный поиск только для физических
            # аптек, у которых цена пока не найдена.
            missing_physical_ids = (
                physical_ids - set(price_map)
            )

            if missing_physical_ids:
                print(
                    f"      Районный поиск: цены пока не найдены "
                    f"для {len(missing_physical_ids)} "
                    f"физических аптек."
                )

                for raon_id, district_ids in sorted(
                    pharmacies_by_raon.items()
                ):
                    target_ids = (
                        missing_physical_ids & district_ids
                    )

                    if not target_ids:
                        continue

                    district_name = next(
                        (
                            str(ph.get("district", "")).strip()
                            for ph in nearby
                            if normalize_id(ph["id"])
                            in target_ids
                        ),
                        f"raon={raon_id}",
                    )

                    print(
                        f"      Ищем в районе "
                        f"{district_name}: "
                        f"нужные ID {sorted(target_ids)}"
                    )

                    try:
                        district_items = search_drug_prices(
                            norm,
                            needed_apteka_ids=target_ids,
                            session=session,
                            raon_id=raon_id,
                        )
                    except Exception as e:
                        print(
                            f"      ОШИБКА районного поиска "
                            f"raon={raon_id}: {e}"
                        )
                        continue

                    update_price_map(
                        price_map,
                        district_items,
                        target_ids,
                    )

                    # После каждого района пересчитываем
                    # недостающие ID: найденные больше не ищем.
                    missing_physical_ids = (
                        physical_ids - set(price_map)
                    )

            physical_found = sum(
                1
                for ph in nearby
                if normalize_id(ph["id"]) in price_map
            )

            internet_found = sum(
                1
                for ph in internet_pharmacies
                if normalize_id(ph["id"]) in price_map
            )

            print(
                f"      Физических аптек с ценой: "
                f"{physical_found}/{len(nearby)}"
            )
            print(
                f"      Интернет-аптек с ценой: "
                f"{internet_found}/{len(internet_pharmacies)}"
            )

            results[raw_name] = price_map

            if config.REQUEST_DELAY:
                time.sleep(config.REQUEST_DELAY)

    finally:
        session.close()

    print(f"\n[5/5] Формируем отчёт {config.REPORT_FILE}...")

    build_report(
        raw_products,
        nearby,
        internet_pharmacies,
        results,
        config.REPORT_FILE,
    )

    print(f"\nГотово! Отчёт сохранён: {config.REPORT_FILE}")


if __name__ == "__main__":
    main()